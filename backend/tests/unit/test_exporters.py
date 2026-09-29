import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.core.dataframe_ops.exporters import export_dataframe
from app.core.errors import DataCleanError


def frame():
    return pd.DataFrame(
        {
            "name": ["Jo", "Al", None],
            "mail": ["a@x.com", "=SUM(A1:A2)", "+CMD|' /C calc'!A0"],
            "n": ["1", "2", "3"],
        }
    )


class TestCsvExport:
    def test_contents(self, tmp_path: Path):
        out = tmp_path / "out.csv"
        stats = export_dataframe(frame(), "csv", out, guard_formulas=False)
        text = out.read_text(encoding="utf-8-sig")
        assert text.splitlines()[0] == "name,mail,n"
        assert stats["rows"] == 3
        assert stats["columns"] == 3

    def test_formula_guard(self, tmp_path: Path):
        out = tmp_path / "out.csv"
        stats = export_dataframe(frame(), "csv", out)
        text = out.read_text(encoding="utf-8-sig")
        assert "'=SUM(A1:A2)" in text
        assert "+CMD|" in text
        assert stats["formula_cells_sanitized"] == 2

    def test_guard_can_be_disabled(self, tmp_path: Path):
        out = tmp_path / "out.csv"
        stats = export_dataframe(frame(), "csv", out, guard_formulas=False)
        assert "=SUM(A1:A2)" in out.read_text(encoding="utf-8-sig")
        assert stats["formula_cells_sanitized"] == 0

    def test_negative_numbers_not_guarded(self, tmp_path: Path):
        out = tmp_path / "out.csv"
        df = pd.DataFrame({"v": ["-5", "+3.5"]})
        stats = export_dataframe(df, "csv", out)
        lines = out.read_text(encoding="utf-8-sig").splitlines()
        assert lines[1:] == ["-5", "+3.5"]
        assert stats["formula_cells_sanitized"] == 0


class TestJsonExport:
    def test_records_shape_and_nulls(self, tmp_path: Path):
        out = tmp_path / "out.json"
        export_dataframe(frame(), "json", out)
        records = json.loads(out.read_text(encoding="utf-8"))
        assert len(records) == 3
        assert records[2]["name"] is None
        assert records[1]["mail"] == "=SUM(A1:A2)"  # JSON needs no formula guard


class TestXlsxExport:
    def test_roundtrip(self, tmp_path: Path):
        out = tmp_path / "out.xlsx"
        stats = export_dataframe(frame(), "xlsx", out)
        book = load_workbook(out)
        sheet = book["Data"]
        assert [c.value for c in sheet[1]] == ["name", "mail", "n"]
        assert sheet.cell(row=2, column=1).value == "Jo"
        assert sheet.freeze_panes == "A2"
        assert sheet.cell(row=1, column=1).font.bold
        assert stats["formula_cells_sanitized"] == 2
        assert sheet.cell(row=3, column=2).value == "'=SUM(A1:A2)"

    def test_unwritable_path(self, tmp_path: Path):
        out = tmp_path / "no" / "such" / "dir" / "out.xlsx"
        with pytest.raises(DataCleanError) as excinfo:
            export_dataframe(frame(), "xlsx", out)
        assert excinfo.value.code == "export_failure"

    def test_unsupported_format(self, tmp_path: Path):
        with pytest.raises(DataCleanError):
            export_dataframe(frame(), "parquet", tmp_path / "out.parquet")
