import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import make_csv, upload_bytes

FORMULA_ROW = "Formula Guy,=SUM(A1:A2),28"


@pytest.fixture()
def job(client: TestClient) -> str:
    return upload_bytes(
        client,
        make_csv(
            FORMULA_ROW,
            "Jane,jane@example.com,28",
            "Jane,jane@example.com,28",
            header="name,mail,age",
        ),
        "exportme.csv",
    )["job_id"]


class TestExportEndpoint:
    def test_export_csv_default_name(self, client: TestClient, job: str):
        response = client.post(f"/api/jobs/{job}/export", json={})
        assert response.status_code == 200
        assert "text/csv" in response.headers["content-type"]
        assert response.headers["content-disposition"].endswith('dataclean_export.csv"')

    def test_export_csv_content_and_guard(self, client: TestClient, job: str):
        response = client.post(
            f"/api/jobs/{job}/export",
            json={"file_format": "csv", "filename": "clean people", "guard_formulas": True},
        )
        disposition = response.headers["content-disposition"]
        assert "clean_people.csv" in disposition
        text = response.content.decode("utf-8-sig")
        assert "'=SUM(A1:A2)" in text  # formula neutralized
        assert text.splitlines()[-1] == "Jane,jane@example.com,28"

    def test_export_csv_without_guard_keeps_formula(self, client: TestClient, job: str):
        response = client.post(
            f"/api/jobs/{job}/export",
            json={"file_format": "csv", "guard_formulas": False},
        )
        assert "=SUM(A1:A2)" in response.content.decode("utf-8-sig")

    def test_export_json(self, client: TestClient, job: str):
        response = client.post(
            f"/api/jobs/{job}/export", json={"file_format": "json", "filename": "out"}
        )
        records = json.loads(response.content.decode("utf-8"))
        assert len(records) == 3
        assert records[0]["name"] == "Formula Guy"
        assert records[0]["mail"] == "=SUM(A1:A2)"

    def test_export_xlsx(self, client: TestClient, job: str):
        import io

        response = client.post(
            f"/api/jobs/{job}/export", json={"file_format": "xlsx", "filename": "out"}
        )
        assert "spreadsheetml" in response.headers["content-type"]
        book = load_workbook(io.BytesIO(response.content))
        sheet = book["Data"]
        assert [c.value for c in sheet[1]] == ["name", "mail", "age"]
        assert sheet.cell(row=2, column=2).value == "'=SUM(A1:A2)"

    def test_export_reflects_transformed_data(self, client: TestClient, job: str):
        client.post(
            f"/api/jobs/{job}/transform",
            json={"steps": [{"id": "d", "type": "dedupe", "config": {"mode": "exact"}}]},
        )
        response = client.post(f"/api/jobs/{job}/export", json={"file_format": "json"})
        records = json.loads(response.content.decode("utf-8"))
        assert len(records) == 2

    def test_export_unknown_format_rejected(self, client: TestClient, job: str):
        response = client.post(f"/api/jobs/{job}/export", json={"file_format": "parquet"})
        assert response.status_code == 422

    def test_export_filename_sanitized(self, client: TestClient, job: str):
        response = client.post(
            f"/api/jobs/{job}/export",
            json={"file_format": "csv", "filename": "../../evil name*.csv"},
        )
        assert "../" not in response.headers["content-disposition"]
        assert "evil_name_.csv" in response.headers["content-disposition"]

    def test_export_job_not_found(self, client: TestClient):
        assert client.post("/api/jobs/missing-job-01/export", json={}).status_code == 404
